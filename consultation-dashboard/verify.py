# -*- coding: utf-8 -*-
"""독립 검증 스크립트.
 ① pandas로 '처음부터' 같은 규칙을 다시 구현해 answer_key.json과 비교  (dashboard.html의 JS 코드를 보지 않음)
 ② dashboard.html을 헤드리스 브라우저(인터넷 차단 상태)로 열어 같은 샘플을 넣고 결과를 answer_key와 비교
    (②는 playwright가 있을 때만: pip install playwright)
사용법:  python verify.py"""
import json, sys, base64, pathlib
from datetime import date, datetime, timedelta
import pandas as pd

ROOT = pathlib.Path(__file__).parent
KEY = json.load(open(ROOT / "test-data" / "answer_key.json", encoding="utf-8"))
EPOCH = pd.Timestamp(1899, 12, 30)

# ---------------- ① pandas 독립 재계산 ----------------
def to_date(v):
    """셀 하나를 date로. 실패하면 'empty' 또는 'error'."""
    if v is None or (isinstance(v, float) and pd.isna(v)) or (isinstance(v, str) and not v.strip()):
        return "empty"
    if isinstance(v, (datetime, pd.Timestamp)): return v.date()
    if isinstance(v, (int, float)):
        return (EPOCH + pd.Timedelta(days=int(v))).date()   # 엑셀 시리얼
    import re
    m = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", str(v))
    if m:
        try: return date(int(m[1]), int(m[2]), int(m[3]))
        except ValueError: pass
    return "error"

def load(path):
    """헤더 행을 찾고(상위 20행) 데이터프레임으로 만든다. 시트가 여럿이면 헤더가 있는 시트."""
    for sheet, raw in pd.read_excel(path, sheet_name=None, header=None, dtype=object).items():
        for i in range(min(20, len(raw))):
            cells = [str(x).replace(" ", "") for x in raw.iloc[i].tolist() if pd.notna(x)]
            if any(("상담일" in c or c == "일자") for c in cells) and any("환자" in c for c in cells):
                df = raw.iloc[i + 1:].copy(); df.columns = [str(x).replace(" ", "") if pd.notna(x) else f"c{j}" for j, x in enumerate(raw.iloc[i])]
                return df.dropna(how="all")
    raise RuntimeError("헤더를 찾지 못함: " + str(path))

def compute(path, meeting, _ly=True):
    paths = path if isinstance(path, (list, tuple)) else [path]
    df = pd.concat([load(p) for p in paths], ignore_index=True)
    datecol = next(c for c in df.columns if "상담일" in c)
    nocol = next((c for c in df.columns if "환자번호" in c), None)
    namecol = next((c for c in df.columns if c == "환자명"), None)
    staffcol = next((c for c in df.columns if "담당" in c), None)
    typecol = next((c for c in df.columns if "유형" in c), None)
    W = datetime.strptime(meeting, "%Y-%m-%d").date()
    assert W.weekday() == 2, "회의일은 수요일이어야 함"
    this_s, this_e = W - timedelta(7), W - timedelta(1)
    prev_s, prev_e = W - timedelta(14), W - timedelta(8)
    Y = this_e.year
    jan1 = date(Y, 1, 1)
    first_wed = jan1 + timedelta((2 - jan1.weekday()) % 7)           # 파이썬: 월=0, 수=2
    week_of = lambda d: 0 if d < first_wed else (d - first_wed).days // 7 + 1

    cnt = dict(loaded=len(df), sum=0, no_date=0, date_error=0, prev_year=0, after=0, counted=0, no_patient_no=0)
    recs = []; allv = []
    for _, r in df.iterrows():
        cells = [str(x).replace(" ", "") for x in r.tolist() if pd.notna(x)]
        if any(c in ("합계", "소계", "총계") for c in cells): cnt["sum"] += 1; continue
        d = to_date(r[datecol])
        if d == "empty": cnt["no_date"] += 1; continue
        if d == "error": cnt["date_error"] += 1; continue
        no = str(r[nocol]).strip() if nocol and pd.notna(r[nocol]) and str(r[nocol]).strip() else ""
        nm = str(r[namecol]).strip() if namecol and pd.notna(r[namecol]) else ""
        allv.append(dict(d=d, no=no, nm=nm))          # 날짜가 정상인 모든 행(연도 무관) = 중복의심 대상
        if d < jan1: cnt["prev_year"] += 1; continue
        if d > this_e: cnt["after"] += 1; continue
        cnt["counted"] += 1
        if not no: cnt["no_patient_no"] += 1
        recs.append(dict(month=d.month, week=week_of(d), pid=("N:" + no) if no else "M:" + nm, d=d,
                         type=(str(r[typecol]).strip() if typecol and pd.notna(r[typecol]) and str(r[typecol]).strip() else "(미분류)"),
                         staff=(str(r[staffcol]).strip() if staffcol and pd.notna(r[staffcol]) and str(r[staffcol]).strip() else "(미지정)")))
    t = pd.DataFrame(recs, columns=["month", "week", "pid", "d", "staff", "type"])
    per = t.groupby("week").agg(patients=("pid", "nunique"), consults=("pid", "size"))   # 주별 중복제거 / 행 수
    kt = week_of(this_e)
    get = lambda k: (int(per.loc[k, "patients"]), int(per.loc[k, "consults"])) if k in per.index else (0, 0)
    ytd_p = int(per.loc[per.index <= kt, "patients"].sum()); ytd_c = int(per.loc[per.index <= kt, "consults"].sum())
    def group_stats(col):                                                  # 직원/유형: 같은 규칙을 그룹 단위로
        res = {}
        for name, g in t.groupby(col):
            sp = g.groupby("week").agg(p=("pid", "nunique"), c=("pid", "size"))
            sg = lambda k: [int(sp.loc[k, "p"]), int(sp.loc[k, "c"])] if k in sp.index else [0, 0]
            res[name] = {"this": sg(kt), "prev": sg(kt - 1),
                         "ytd": [int(sp.loc[sp.index <= kt, "p"].sum()), int(sp.loc[sp.index <= kt, "c"].sum())]}
        return res
    staff = group_stats("staff"); types = group_stats("type")
    mon = t.groupby("month").agg(p=("pid", "nunique"), c=("pid", "size"))   # 월간: 달력 월 안에서 중복제거
    months = {str(k): [int(v.p), int(v.c)] for k, v in mon.iterrows()}
    av = pd.DataFrame(allv, columns=["d", "no", "nm"])
    both = av[(av.no != "") & (av.nm != "")]
    dupes = {"A": int((both.groupby("no").nm.nunique() > 1).sum()),                   # 같은 번호·다른 이름
             "B": int((both.groupby("nm").no.nunique() > 1).sum()),                   # 같은 이름·다른 번호
             "C": int((av.assign(k=av.apply(lambda r: ("N:" + r.no) if r.no else ("M:" + r.nm if r.nm else None), axis=1)).dropna(subset=["k"]).groupby(["d", "k"]).size() > 1).sum())}   # 같은 날 같은 환자 2회+
    last_year = None
    if _ly:                                                   # 전년 동기 = 같은 계산을 52주(364일) 전 회의일로
        r = compute(path, str(W - timedelta(364)), _ly=False)
        if r["counts"]["counted"] > 0: last_year = {k2: r[k2] for k2 in ("this", "prev", "ytd")}
    return {"types": types, "dupes": dupes, "last_year": last_year, "months": months, "staff": staff, "this": dict(zip(("patients", "consults"), get(kt))), "prev": dict(zip(("patients", "consults"), get(kt - 1))),
            "ytd": {"patients": ytd_p, "consults": ytd_c},
            "weeks_nonzero": {str(k): [int(v.patients), int(v.consults)] for k, v in per.iterrows()}, "counts": cnt}

# ---------------- 비교/출력 ----------------
fails = 0; total = 0
def check(label, got, want):
    global fails, total
    total += 1
    ok = got == want
    if not ok: fails += 1
    print(("  PASS " if ok else "  FAIL ") + label + ("" if ok else f"\n        기대={want}\n        실제={got}"))

def compare(tag, got, exp):
    for k in ("this", "prev", "ytd"): check(f"{tag} {k}", got[k], exp[k])
    if "weeks_nonzero" in exp: check(f"{tag} 주별표(0 아닌 주)", got["weeks_nonzero"], exp["weeks_nonzero"])
    if "counts" in exp: check(f"{tag} 품질 카운트", got["counts"], exp["counts"])
    if "dupes_exp" in exp: check(f"{tag} 중복의심 건수", got["dupes"], exp["dupes_exp"])
    if "last_year" in exp: check(f"{tag} 전년 동기", got.get("last_year", got.get("lastYear")), exp["last_year"])
    if "months" in exp: check(f"{tag} 월간", got["months"], exp["months"])
    if "types" in exp: check(f"{tag} 상담유형별", got["types"], exp["types"])
    if "staff" in exp: check(f"{tag} 직원별", got["staff"], exp["staff"])

print("=== ① pandas 독립 재계산  vs  answer_key ===")
for f, spec in KEY["samples"].items():
    print(f"[{f}]")
    for sc in spec["scenarios"]:
        compare(f"회의일 {sc['meeting']}", compute(ROOT / "test-data" / f, sc["meeting"]), sc["expect"])
    if f in KEY["dupes"]: check("중복의심 건수(A번호-이름/B이름-번호/C같은날)", compute(ROOT / "test-data" / f, spec["scenarios"][0]["meeting"])["dupes"], KEY["dupes"][f])

print("[병합: 파일 여러 개]")
for sc in KEY["multi"]:
    compare("병합 " + "+".join(sc["files"]), compute([ROOT / "test-data" / f for f in sc["files"]], sc["meeting"]), sc["expect"])

print("\n=== ② dashboard.html (헤드리스 브라우저, 인터넷 차단)  vs  answer_key ===")
try:
    from playwright.sync_api import sync_playwright
except ImportError:
    print("  SKIP playwright 미설치 (pip install playwright) — ②는 건너뜀"); sync_playwright = None
if sync_playwright:
    exe = next(iter(sorted(pathlib.Path("/opt/pw-browsers").glob("chromium-*/chrome-linux/chrome"))), None)
    with sync_playwright() as p:
        b = p.chromium.launch(**({"executable_path": str(exe)} if exe else {}))
        ctx = b.new_context(); ctx.set_offline(True)            # T10: 인터넷 차단
        reqs = []; ctx.on("request", lambda r: reqs.append(r.url) if not r.url.startswith(("file:", "data:", "blob:")) else None)
        page = ctx.new_page(); errs = []; page.on("pageerror", lambda e: errs.append(str(e)))
        page.goto((ROOT / "dashboard.html").resolve().as_uri())
        def upload(sel, files):                 # 파일 넣기 → 대시보드가 '처리 완료' 신호를 올릴 때까지 대기 (경주 조건 방지)
            n = int(page.evaluate("document.documentElement.dataset.loads || 0"))
            page.set_input_files(sel, files)
            page.wait_for_function("n => +(document.documentElement.dataset.loads || 0) > n", arg=n)
        check("T10 외부 요청 0건", reqs, []); check("페이지 JS 오류 없음", errs, [])
        check("내장 자가진단 전부 통과", page.inner_text("#foot").split("통과")[0].strip().replace("내장 자가진단 ", "").split("/")[0] ==
              page.inner_text("#foot").split("통과")[0].strip().split("/")[1], True)
        for f, spec in KEY["samples"].items():
            print(f"[{f}]")
            b64 = base64.b64encode((ROOT / "test-data" / f).read_bytes()).decode()
            for sc in spec["scenarios"]:
                got = page.evaluate("""([b64, m]) => { const bin = atob(b64); const u = new Uint8Array(bin.length);
                    for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i);
                    return window.DashboardAPI.analyzeWorkbook(u.buffer, m); }""", [b64, sc["meeting"]])
                compare(f"회의일 {sc['meeting']}", got, sc["expect"])
            if f in KEY["dupes"]: check("중복의심 건수(브라우저)", got["dupes"], KEY["dupes"][f])
        # 실제 화면(업로드 → KPI 숫자) 확인
        upload("#fileInput", str(ROOT / "test-data" / "sample1_normal.xlsx"))
        page.fill("#meetingDate", "2026-10-07"); page.dispatch_event("#meetingDate", "change")
        nums = page.eval_on_selector_all(".kpi .num", "els => els.map(e => parseInt(e.textContent))")
        check("화면 KPI 4개 (환자3·상담5·YTD 5·8)", nums, [3, 5, 5, 8])
        check("화면 주간 표 행 존재", page.locator("#weekTable tr.cur").count(), 1)
        # --- 인쇄/PDF(Phase 3): 인쇄 화면에서 조작 버튼은 숨고 핵심 내용은 보이는지, PDF가 만들어지는지 ---
        upload("#fileInput", str(ROOT / "test-data" / "sample1_normal.xlsx")); setm0 = page.fill("#meetingDate", "2026-10-07"); page.dispatch_event("#meetingDate", "change")
        page.emulate_media(media="print")
        vis = {k: page.is_visible(k) for k in ["#printHead", "#kpis", "#chartBox", "#weekTable", "#printBtn", "#dropzone", "#meetingDate"]}
        check("인쇄 화면: 머리글·KPI·차트·표 보임 / 버튼·업로드·회의일 입력 숨김", vis,
              {"#printHead": True, "#kpis": True, "#chartBox": True, "#weekTable": True, "#printBtn": False, "#dropzone": False, "#meetingDate": False})
        pdf = page.pdf(prefer_css_page_size=True, print_background=True)
        check("PDF 생성(%PDF, 내용 있음)", pdf[:5] == b"%PDF-" and len(pdf) > 20000, True)
        page.emulate_media(media="screen")
        # --- 다중 파일 병합(Phase 3): API로 값 비교 + 화면에서 '파일 추가' 후 전년 동기 표시 ---
        print("[병합: 파일 여러 개]")
        for sc in KEY["multi"]:
            bs = [base64.b64encode((ROOT / "test-data" / f).read_bytes()).decode() for f in sc["files"]]
            got = page.evaluate("""([bs, m]) => { const bufs = bs.map((b64) => { const bin = atob(b64); const u = new Uint8Array(bin.length);
                for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i); return u.buffer; });
                return window.DashboardAPI.analyzeFiles(bufs, m); }""", [bs, sc["meeting"]])
            compare("병합(브라우저) " + "+".join(sc["files"]), got, sc["expect"])
        upload("#fileInput", str(ROOT / "test-data" / "sample1_normal.xlsx")); page.fill("#meetingDate", "2026-10-07"); page.dispatch_event("#meetingDate", "change")
        check("병합 화면① 파일 1개면 전년 동기 줄 없음", page.locator(".kpi .ly").count(), 0)
        upload("#fileInputAdd", str(ROOT / "test-data" / "sample5_2025.xlsx"))
        check("병합 화면② 파일 추가 후 2026 KPI 그대로(3·5·5·8)", page.eval_on_selector_all(".kpi .num", "els => els.map(e => parseInt(e.textContent))"), [3, 5, 5, 8])
        check("병합 화면③ 전년 동기 값 표시(3명·3건·5명·5건)", page.eval_on_selector_all(".kpi .ly", "els => els.map(e => e.textContent.match(/전년 동기 (\\d+)/)[1])"), ["3", "3", "5", "5"])
        check("병합 화면④ 파일 칩 2개", page.locator("#fileList .chip").count(), 2)
        upload("#fileInputAdd", str(ROOT / "test-data" / "sample2_irregular.xlsx"))
        check("병합 화면⑤ 날짜 겹치는 파일 → 경고", "겹칩니다" in page.inner_text("#warnBanner"), True)
        page.click("#fileList .chip:last-child .x")
        check("병합 화면⑥ ✕로 파일 빼기", page.locator("#fileList .chip").count(), 2)
        # --- 상담유형 카드: 유형 열이 있으면 보이고(초진·재진 행), 없으면 숨김 ---
        upload("#fileInput", str(ROOT / "test-data" / "sample1_normal.xlsx")); page.fill("#meetingDate", "2026-10-07"); page.dispatch_event("#meetingDate", "change")
        check("상담유형 화면: 카드 보임 + 초진·재진 행", [page.is_visible("#typeCard"), "초진" in page.inner_text("#typeTable"), "재진" in page.inner_text("#typeTable")], [True, True, True])
        upload("#fileInput", str(ROOT / "test-data" / "sample3_edge.xlsx"))
        check("상담유형 화면: 유형 열이 없으면 카드 숨김", page.is_hidden("#typeCard"), True)
        # --- 중복의심 리포트 화면: sample6 업로드 → 3건(유형별 1건) 표시 ---
        upload("#fileInput", str(ROOT / "test-data" / "sample6_dupes.xlsx"))
        page.click("#dupeSummary")                    # 접혀 있는 패널을 펼쳐야 내용이 보임
        check("중복의심 화면: 제목 '3건'", "(3건)" in page.inner_text("#dupeSummary"), True)
        check("중복의심 화면: 유형 3종 표시", all(t in page.inner_text("#dupeTable") for t in ("같은 번호·다른 이름", "같은 이름·다른 번호", "같은 날 여러 번 상담")), True)
        upload("#fileInput", str(ROOT / "test-data" / "sample1_normal.xlsx")); page.fill("#meetingDate", "2026-10-07"); page.dispatch_event("#meetingDate", "change")
        upload("#fileInputAdd", str(ROOT / "test-data" / "sample5_2025.xlsx"))   # 전년 동기 줄이 엑셀에 들어가는지 보려고 2025 파일 추가
        # --- 결과 엑셀 다운로드(Phase 3): 받은 파일을 openpyxl로 열어 정답지와 비교 ---
        import openpyxl, tempfile
        with page.expect_download() as dl: page.click("#downloadBtn")
        out = pathlib.Path(tempfile.gettempdir()) / "verify_result.xlsx"; dl.value.save_as(out)
        wbx = openpyxl.load_workbook(out, data_only=True)
        rows = {r[0]: r[1:] for r in wbx["요약"].iter_rows(values_only=True) if r and r[0]}
        check("엑셀 저장: 전년 동기 줄(3/3·YTD 5/5)", [list(rows["전년 동기 환자수"][:1]) + [rows["전년 동기 환자수"][3]], list(rows["전년 동기 상담건수"][:1]) + [rows["전년 동기 상담건수"][3]]], [[3, 5], [3, 5]])
        e0 = KEY["samples"]["sample1_normal.xlsx"]["scenarios"][0]["expect"]
        d = lambda k: e0["this"][k] - e0["prev"][k]
        check("엑셀 저장: 요약 시트 이번주·전주·증감", [list(rows["환자수"][:3]), list(rows["상담건수"][:3])],
              [[e0["this"]["patients"], e0["prev"]["patients"], d("patients")], [e0["this"]["consults"], e0["prev"]["consults"], d("consults")]])
        check("엑셀 저장: YTD 환자/상담", [rows["환자수"][3], rows["상담건수"][3]], [e0["ytd"]["patients"], e0["ytd"]["consults"]])
        wk = {r[0]: r[3:5] for r in wbx["주간표"].iter_rows(min_row=2, values_only=True)}
        check("엑셀 저장: 주간표(0 아닌 주)", {k[1:]: list(v) for k, v in wk.items() if v[0] or v[1]}, e0["weeks_nonzero"])
        check("엑셀 저장: 시트 구성", wbx.sheetnames, ["요약", "주간표", "직원별", "상담유형별", "중복의심", "월간", "제외된 행"])
        # --- 컬럼 매핑 기억(Phase 2): 자동 감지 실패 → 수동 지정 → 새로고침 후 자동 적용 → 지우기 ---
        s4 = str(ROOT / "test-data" / "sample4_custom_headers.xlsx")
        kpis = lambda: page.eval_on_selector_all(".kpi .num", "els => els.map(e => parseInt(e.textContent))")
        setm = lambda: (page.fill("#meetingDate", "2026-10-07"), page.dispatch_event("#meetingDate", "change"))
        upload("#fileInput", s4)
        check("기억① 영문 제목은 자동 감지 실패(결과 숨김, 지정창 열림)", [page.is_hidden("#result"), page.evaluate("document.querySelector('#mapPanel').open")], [True, True])
        for sel, v in (("#mapDate", "0"), ("#mapNo", "1"), ("#mapName", "2"), ("#mapStaff", "3")): page.select_option(sel, v)
        page.click("#applyMapBtn"); setm()
        check("기억② 수동 지정 후 계산 (환자3·상담5·YTD 5·8)", kpis(), [3, 5, 5, 8])
        page.reload(); upload("#fileInput", s4); setm()
        check("기억③ 새로고침 후 같은 양식 자동 적용", [page.is_visible("#result"), kpis()], [True, [3, 5, 5, 8]])
        page.click("#mapPanel summary"); page.click("#clearMemBtn")
        upload("#fileInput", s4)
        check("기억④ 설정 지우면 다시 자동 감지 실패", page.is_hidden("#result"), True)
        b.close()

# ---------------- ③ 시연 자료(가상 1,300여 행) 교차검증: 정답지 없이 '서로 독립인 두 구현'끼리 비교 ----------------
DEMO = ROOT / "test-data" / "demo_hospital_2026.xlsx"
if DEMO.exists() and sync_playwright:
    print("\n=== ③ 시연 자료 교차검증 (pandas  vs  대시보드, 정답지 없이 서로 비교) ===")
    b64 = base64.b64encode(DEMO.read_bytes()).decode()
    with sync_playwright() as p:
        b = p.chromium.launch(**({"executable_path": str(exe)} if exe else {})); pg = b.new_page()
        pg.goto((ROOT / "dashboard.html").resolve().as_uri())
        for m in ["2026-01-07", "2026-01-14", "2026-03-04", "2026-06-17", "2026-09-30", "2026-10-07"]:
            got = pg.evaluate("""([b64, m]) => { const bin = atob(b64); const u = new Uint8Array(bin.length);
                for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i); return window.DashboardAPI.analyzeWorkbook(u.buffer, m); }""", [b64, m])
            exp = compute(DEMO, m)
            for k in ("this", "prev", "ytd", "weeks_nonzero", "counts", "staff", "types", "months", "dupes"):
                check(f"시연자료 {m} {k}", got[k], exp[k])
        b.close()

print(f"\n결과: {total - fails}/{total} PASS" + ("  ✅ 전부 통과" if fails == 0 else f"  ❌ {fails}개 FAIL"))
sys.exit(1 if fails else 0)
